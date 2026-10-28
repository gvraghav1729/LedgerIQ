
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

from financial_metrics import calculate_financial_metrics
from anomaly_detector import detect_anomalies
from report_generator import (
    generate_excel_report,
    generate_pdf_report,
)


def make_sample_data():
    return pd.DataFrame([
        {
            "period": "2023",
            "revenue": 1000,
            "cost_of_sales": 600,
            "operating_expenses": 200,
            "total_assets": 1500,
            "current_assets": 500,
            "current_liabilities": 250,
            "total_liabilities": 700,
            "equity": 800,
            "net_income": 100,
        },
        {
            "period": "2024",
            "revenue": 850,
            "cost_of_sales": 550,
            "operating_expenses": 220,
            "total_assets": 1600,
            "current_assets": 520,
            "current_liabilities": 260,
            "total_liabilities": 730,
            "equity": 870,
            "net_income": 60,
        },
    ])


def test_excel_report_contains_expected_sheets():
    data = make_sample_data()
    metrics = calculate_financial_metrics(data)
    findings = detect_anomalies(data)

    excel_bytes = generate_excel_report(metrics, findings)

    workbook = load_workbook(BytesIO(excel_bytes), read_only=True)

    assert "Financial Analysis" in workbook.sheetnames
    assert "Anomaly Findings" in workbook.sheetnames

    workbook.close()


def test_excel_report_contains_financial_data():
    data = make_sample_data()
    metrics = calculate_financial_metrics(data)
    findings = detect_anomalies(data)

    excel_bytes = generate_excel_report(metrics, findings)

    workbook = load_workbook(BytesIO(excel_bytes), read_only=True)
    sheet = workbook["Financial Analysis"]

    assert sheet["A2"].value == "2023"
    assert sheet["B2"].value == 1000

    workbook.close()


def test_pdf_report_is_generated():
    data = make_sample_data()
    metrics = calculate_financial_metrics(data)
    findings = detect_anomalies(data)

    pdf_bytes = generate_pdf_report(metrics, findings)

    assert pdf_bytes.startswith(b"%PDF")

    reader = PdfReader(BytesIO(pdf_bytes))
    assert len(reader.pages) >= 1


def test_pdf_report_contains_report_title():
    data = make_sample_data()
    metrics = calculate_financial_metrics(data)
    findings = detect_anomalies(data)

    pdf_bytes = generate_pdf_report(metrics, findings)

    reader = PdfReader(BytesIO(pdf_bytes))
    text = "\n".join(
        page.extract_text() or ""
        for page in reader.pages
    )

    assert "LedgerIQ Financial Analysis Report" in text
    assert "Unusual Financial Movements" in text
