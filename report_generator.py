
from io import BytesIO

import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    LongTable,
    TableStyle,
    PageBreak,
)
from xml.sax.saxutils import escape


# --------------------------------------------------
# 1. Formatting helpers
# --------------------------------------------------

def format_inr(value):
    """Format an amount using Indian comma grouping and the INR symbol."""

    if pd.isna(value):
        return "N/A"

    amount = int(round(float(value)))
    sign = "-" if amount < 0 else ""
    digits = str(abs(amount))

    if len(digits) <= 3:
        formatted = digits
    else:
        last_three = digits[-3:]
        remaining = digits[:-3]
        groups = []

        while remaining:
            groups.insert(0, remaining[-2:])
            remaining = remaining[:-2]

        formatted = ",".join(groups) + "," + last_three

    return f"{sign}₹{formatted}"

def register_report_fonts():
    """Register Windows Arial fonts with support for the rupee symbol."""

    regular_path = r"C:\Windows\Fonts\arial.ttf"
    bold_path = r"C:\Windows\Fonts\arialbd.ttf"

    try:
        pdfmetrics.registerFont(
            TTFont("LedgerIQSans", regular_path)
        )
        pdfmetrics.registerFont(
            TTFont("LedgerIQSans-Bold", bold_path)
        )
        return "LedgerIQSans", "LedgerIQSans-Bold"

    except (OSError, Exception):
        # Fallback if Arial is unavailable.
        return "Helvetica", "Helvetica-Bold"

# --------------------------------------------------
# 2. Excel report
# --------------------------------------------------

def generate_excel_report(metrics, findings):
    """Generate an Excel workbook with analysis and anomaly findings."""

    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        metrics.to_excel(
            writer,
            sheet_name="Financial Analysis",
            index=False,
        )

        findings.to_excel(
            writer,
            sheet_name="Anomaly Findings",
            index=False,
        )

        workbook = writer.book

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="17365D",
        )

        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions

            for cell in worksheet[1]:
                cell.font = Font(
                    name="DejaVu Sans",
                    bold=True,
                    color="FFFFFF",
                )
                cell.fill = header_fill
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

            for column_cells in worksheet.columns:
                column_letter = column_cells[0].column_letter
                max_length = max(
                    len(str(cell.value or ""))
                    for cell in column_cells
                )
                worksheet.column_dimensions[column_letter].width = min(
                    max(max_length + 2, 12),
                    45,
                )

    output.seek(0)
    return output.getvalue()


# --------------------------------------------------
# 3. PDF report
# --------------------------------------------------

def generate_pdf_report(metrics, findings):
    """
    Create a structured two-page PDF.

    Page 1: Summary and financial metrics.
    Page 2: Unusual financial movements and disclaimer.
    """

    output = BytesIO()

    page_size = landscape(A4)
    page_width, page_height = page_size

    margin = 15 * mm
    available_width = page_width - 2 * margin

    document = SimpleDocTemplate(
        output,
        pagesize=page_size,
        rightMargin=margin,
        leftMargin=margin,
        topMargin=13 * mm,
        bottomMargin=13 * mm,
        title="LedgerIQ Financial Analysis Report",
        author="LedgerIQ",
    )

    regular_font, bold_font = register_report_fonts()

    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name="LedgerTitle",
            parent=styles["Title"],
            fontName=bold_font,
            fontSize=24,
            leading=29,
            alignment=TA_CENTER,
            spaceAfter=7 * mm,
            textColor=colors.HexColor("#17365D"),
        )
    )

    styles.add(
        ParagraphStyle(
            name="LedgerHeading",
            parent=styles["Heading2"],
            fontName=bold_font,
            fontSize=16,
            leading=20,
            spaceBefore=4 * mm,
            spaceAfter=4 * mm,
            textColor=colors.HexColor("#17365D"),
        )
    )

    styles.add(
        ParagraphStyle(
            name="LedgerBody",
            parent=styles["BodyText"],
            fontName=regular_font,
            fontSize=11,
            leading=15,
            spaceAfter=3,
        )
    )

    styles.add(
        ParagraphStyle(
            name="LedgerCell",
            parent=styles["BodyText"],
            fontName=regular_font,
            fontSize=10,
            leading=12,
            wordWrap="CJK",
        )
    )

    styles.add(
        ParagraphStyle(
            name="LedgerCellHeader",
            parent=styles["BodyText"],
            fontName=bold_font,
            fontSize=10,
            leading=12,
            textColor=colors.white,
            wordWrap="CJK",
        )
    )

    story = []

    # --------------------------------------------------
    # PAGE 1 — Summary and financial metrics
    # --------------------------------------------------

    story.append(
        Paragraph(
            "LedgerIQ Financial Analysis Report",
            styles["LedgerTitle"],
        )
    )

    latest = metrics.iloc[-1]

    summary_text = (
        f"<b>Reporting period:</b> {escape(str(latest['period']))}<br/>"
        f"<b>Revenue:</b> {format_inr(latest['revenue'])}<br/>"
        f"<b>Net income:</b> {format_inr(latest['net_income'])}<br/>"
        f"<b>Net profit margin:</b> "
        f"{latest['net_profit_margin_pct']:.2f}%"
        if pd.notna(latest["net_profit_margin_pct"])
        else (
            f"<b>Reporting period:</b> {escape(str(latest['period']))}<br/>"
            f"<b>Revenue:</b> {format_inr(latest['revenue'])}<br/>"
            f"<b>Net income:</b> {format_inr(latest['net_income'])}<br/>"
            "<b>Net profit margin:</b> N/A"
        )
    )

    story.append(
        Paragraph(summary_text, styles["LedgerBody"])
    )
    story.append(Spacer(1, 5 * mm))

    story.append(
        Paragraph(
            "Financial Metrics by Period",
            styles["LedgerHeading"],
        )
    )

    metric_columns = [
        "period",
        "revenue",
        "net_income",
        "gross_profit_margin_pct",
        "net_profit_margin_pct",
        "current_ratio",
        "debt_to_equity_ratio",
    ]

    metric_headers = [
        "Period",
        "Revenue (INR)",
        "Net income (INR)",
        "Gross margin (%)",
        "Net margin (%)",
        "Current ratio",
        "Liabilities / equity",
    ]

    metric_data = [[
        Paragraph(escape(header), styles["LedgerCellHeader"])
        for header in metric_headers
    ]]

    for _, row in metrics[metric_columns].iterrows():
        metric_data.append([
            Paragraph(escape(str(row["period"])), styles["LedgerCell"]),
            Paragraph(escape(format_inr(row["revenue"])), styles["LedgerCell"]),
            Paragraph(escape(format_inr(row["net_income"])), styles["LedgerCell"]),
            Paragraph(
                f"{row['gross_profit_margin_pct']:.2f}",
                styles["LedgerCell"],
            ),
            Paragraph(
                f"{row['net_profit_margin_pct']:.2f}"
                if pd.notna(row["net_profit_margin_pct"])
                else "N/A",
                styles["LedgerCell"],
            ),
            Paragraph(
                f"{row['current_ratio']:.2f}"
                if pd.notna(row["current_ratio"])
                else "N/A",
                styles["LedgerCell"],
            ),
            Paragraph(
                f"{row['debt_to_equity_ratio']:.2f}"
                if pd.notna(row["debt_to_equity_ratio"])
                else "N/A",
                styles["LedgerCell"],
            ),
        ])

    metric_widths = [
        available_width * fraction
        for fraction in [0.08, 0.17, 0.17, 0.15, 0.15, 0.12, 0.16]
    ]

    metrics_table = LongTable(
        metric_data,
        colWidths=metric_widths,
        repeatRows=1,
        hAlign="LEFT",
    )

    metrics_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#B8C2CC")),
        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, -1),
            [colors.white, colors.HexColor("#F1F4F8")],
        ),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))

    story.append(metrics_table)

    # Force all anomaly findings to begin on page 2.
    story.append(PageBreak())

    # --------------------------------------------------
    # PAGE 2 — Unusual financial movements
    # --------------------------------------------------

    story.append(
        Paragraph(
            "Unusual Financial Movements",
            styles["LedgerHeading"],
        )
    )

    story.append(
        Paragraph(
            "The following alerts identify unusual movements for further "
            "review. They do not establish that an error or fraud occurred.",
            styles["LedgerBody"],
        )
    )
    story.append(Spacer(1, 5 * mm))

    finding_headers = [
        "Period",
        "Severity",
        "Finding",
        "Explanation",
    ]

    finding_data = [[
        Paragraph(escape(header), styles["LedgerCellHeader"])
        for header in finding_headers
    ]]

    if findings.empty:
        finding_data.append([
            Paragraph("—", styles["LedgerCell"]),
            Paragraph("—", styles["LedgerCell"]),
            Paragraph("No findings", styles["LedgerCell"]),
            Paragraph(
                "No movements crossed the configured screening thresholds.",
                styles["LedgerCell"],
            ),
        ])
    else:
        for _, row in findings.iterrows():
            finding_data.append([
                Paragraph(
                    escape(str(row["period"])),
                    styles["LedgerCell"],
                ),
                Paragraph(
                    escape(str(row["severity"])),
                    styles["LedgerCell"],
                ),
                Paragraph(
                    escape(str(row["finding"])),
                    styles["LedgerCell"],
                ),
                Paragraph(
                    escape(str(row["explanation"])),
                    styles["LedgerCell"],
                ),
            ])

    finding_widths = [
        available_width * fraction
        for fraction in [0.09, 0.11, 0.25, 0.55]
    ]

    findings_table = LongTable(
        finding_data,
        colWidths=finding_widths,
        repeatRows=1,
        hAlign="LEFT",
    )

    findings_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#B8C2CC")),
        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, -1),
            [colors.white, colors.HexColor("#F1F4F8")],
        ),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))

    story.append(findings_table)
    story.append(Spacer(1, 8 * mm))

    story.append(
        Paragraph(
            "<i>Screening alerts are indicators for further review, not "
            "proof of error or fraud. This report is generated by LedgerIQ "
            "and is not a formal audit opinion or investment recommendation.</i>",
            styles["LedgerBody"],
        )
    )

    document.build(story)

    output.seek(0)
    return output.getvalue()
